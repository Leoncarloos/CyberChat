"""Revisión humana de una muestra de respuestas y comparación con el juez de RAGAS.

Dos pasos:

  1. `muestra`: arma una hoja XLSX con 40 respuestas (32 documentales y 8 de seguimiento, una
     repetición por consulta, semilla fija) para que cada revisor la complete POR SEPARADO. La hoja
     no muestra las notas del juez.
  2. `analizar`: lee las dos hojas completadas y calcula el acuerdo entre revisores (kappa ponderado
     cuadrático) y entre cada métrica del juez y la media humana (Spearman).

Escala, igual en las dos preguntas: 0 = no, 1 = en parte, 2 = sí.
  fidelidad:  ¿todo lo que afirma la respuesta está respaldado por los fragmentos recuperados?
  relevancia: ¿la respuesta contesta lo que se preguntó?

Uso:
    python rag-eval/evaluate/human_review.py muestra  --results rag-eval/results/<run_id>.jsonl
    python rag-eval/evaluate/human_review.py analizar --results rag-eval/results/<run_id>.jsonl \
        --revisor-a hoja_a.xlsx --revisor-b hoja_b.xlsx
"""
from __future__ import annotations

import argparse
import json
import random
import statistics as st
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SEMILLA = 20261004
CUPO = {"documental": 32, "seguimiento": 8}
ESCALA = [0, 1, 2]
COLS = ["id", "tipo", "pregunta", "historial", "respuesta", "fragmentos recuperados", "referencia",
        "fidelidad (0/1/2)", "relevancia (0/1/2)", "qué comprobé", "revisor", "fecha"]


def cargar(p: Path) -> list[dict]:
    return [json.loads(l) for l in p.read_text(encoding="utf-8").splitlines() if l.strip()]


def elegibles(results: Path) -> list[dict]:
    """Respuestas que el juez también puntuó: una por consulta, la primera repetición válida."""
    ragas = {(r["query_id"], r["run"]) for r in cargar(results.with_suffix(".ragas.jsonl"))
             if r["metric_scope"] == "generation" and r.get("faithfulness") is not None}
    vistos, out = set(), []
    for r in cargar(results):
        k = (r["query_id"], r["run"])
        if r["dataset"] == "conjunto" and k in ragas and r["query_id"] not in vistos:
            vistos.add(r["query_id"])
            out.append(r)
    return out


def muestra(args) -> None:
    pool = elegibles(args.results)
    rnd = random.Random(SEMILLA)
    elegidas = []
    for tipo, n in CUPO.items():
        g = [r for r in pool if r["tipo"] == tipo]
        elegidas += rnd.sample(g, min(n, len(g)))
    rnd.shuffle(elegidas)

    wb = Workbook()
    ws = wb.active
    ws.title = "Revision"
    ws.append(COLS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
    for r in elegidas:
        hist = "\n".join(f"{m['role']}: {m['content']}" for m in (r.get("history") or []))
        frag = "\n\n".join(f"[{i + 1}] {c['content']}" for i, c in enumerate(r["retrieved_contexts"]))
        ws.append([f"{r['query_id']}-r{r['run']}", r["tipo"], r["user_input"], hist, r["response"],
                   frag, r["reference"], None, None, None, None, None])
    for col, w in zip("ABCDEFGHIJKL", (10, 12, 38, 30, 60, 60, 40, 14, 14, 36, 22, 12)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"H2:I{len(elegidas) + 1}")
    ws.freeze_panes = "C2"

    ins = wb.create_sheet("Instrucciones", 0)
    for linea in [
        "Revisión de respuestas del chatbot (40 respuestas)",
        "",
        "Trabaje por separado, sin consultar al otro revisor. En la hoja «Revision», para cada fila:",
        "1. Lea la pregunta, el historial (si hay), la respuesta y los fragmentos recuperados.",
        "2. FIDELIDAD: 2 = todo lo que afirma la respuesta está en los fragmentos; 1 = hay afirmaciones",
        "   que no están (aunque sean razonables); 0 = contradice los fragmentos o inventa lo central.",
        "   Califique contra los FRAGMENTOS, no contra la referencia ni contra lo que usted sabe.",
        "3. RELEVANCIA: 2 = contesta lo preguntado; 1 = contesta en parte o se desvía; 0 = no contesta.",
        "4. En «qué comprobé» escriba una frase (por ejemplo, qué afirmación no estaba respaldada).",
        "5. Firme con su nombre y la fecha en cada fila.",
        "",
        "La hoja no muestra las notas del juez automático, a propósito.",
    ]:
        ins.append([linea])
    ins["A1"].font = Font(bold=True, size=13)
    ins.column_dimensions["A"].width = 100

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    n_doc = sum(1 for r in elegidas if r["tipo"] == "documental")
    print(f"{len(elegidas)} respuestas ({n_doc} documentales, {len(elegidas) - n_doc} de seguimiento) -> {args.out}")


def kappa_cuadratico(a: list[int], b: list[int]) -> float | None:
    n = len(a)
    k = len(ESCALA)
    obs = [[0.0] * k for _ in range(k)]
    for x, y in zip(a, b):
        obs[x][y] += 1
    ha = [sum(obs[i]) for i in range(k)]
    hb = [sum(obs[i][j] for i in range(k)) for j in range(k)]
    num = den = 0.0
    for i in range(k):
        for j in range(k):
            w = (i - j) ** 2 / (k - 1) ** 2
            num += w * obs[i][j]
            den += w * ha[i] * hb[j] / n
    return None if den == 0 else 1 - num / den


def rangos(v: list[float]) -> list[float]:
    orden = sorted(range(len(v)), key=lambda i: v[i])
    r = [0.0] * len(v)
    i = 0
    while i < len(v):
        j = i
        while j + 1 < len(v) and v[orden[j + 1]] == v[orden[i]]:
            j += 1
        for t in range(i, j + 1):
            r[orden[t]] = (i + j) / 2 + 1
        i = j + 1
    return r


def spearman(x: list[float], y: list[float]) -> float | None:
    if len(x) < 3:
        return None
    rx, ry = rangos(x), rangos(y)
    mx, my = st.mean(rx), st.mean(ry)
    sx = sum((a - mx) ** 2 for a in rx) ** 0.5
    sy = sum((b - my) ** 2 for b in ry) ** 0.5
    return None if sx == 0 or sy == 0 else sum((a - mx) * (b - my) for a, b in zip(rx, ry)) / (sx * sy)


def leer(path: Path) -> dict[str, dict]:
    ws = load_workbook(path)["Revision"]
    h = [c.value for c in ws[1]]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(h, row))
        if d["id"]:
            out[d["id"]] = d
    return out


def analizar(args) -> None:
    A, B = leer(args.revisor_a), leer(args.revisor_b)
    comunes = [i for i in A if i in B]
    ragas = {(r["query_id"], r["run"]): r for r in cargar(args.results.with_suffix(".ragas.jsonl"))
             if r["metric_scope"] == "generation"}
    print(f"Filas comunes: {len(comunes)} · A: {len(A)} · B: {len(B)}")
    for metrica, col, juez in (("fidelidad", "fidelidad (0/1/2)", "faithfulness"),
                               ("relevancia", "relevancia (0/1/2)", "answer_relevancy")):
        ids = [i for i in comunes if A[i][col] in ESCALA and B[i][col] in ESCALA]
        if len(ids) < len(comunes):
            print(f"  {metrica}: {len(comunes) - len(ids)} filas sin calificar por alguno de los dos")
        a = [int(A[i][col]) for i in ids]
        b = [int(B[i][col]) for i in ids]
        k = kappa_cuadratico(a, b)
        exacto = sum(x == y for x, y in zip(a, b)) / len(ids)
        hum = [st.mean([x, y]) for x, y in zip(a, b)]
        jz = []
        for i in ids:
            q, r = i.rsplit("-r", 1)
            jz.append(ragas[(q, int(r))][juez])
        rho = spearman(hum, jz)
        ktxt = "n/d (sin variación)" if k is None else f"{k:.2f}"
        rtxt = "n/d" if rho is None else f"{rho:.2f}"
        print(f"\n{metrica.upper()}  (n={len(ids)})")
        print(f"  acuerdo exacto A-B: {exacto:.2f} · kappa ponderado cuadrático: {ktxt}")
        print(f"  Spearman humano (media A,B) vs juez {juez}: {rtxt}")
        for nivel in ESCALA:
            g = [j for h_, j in zip(hum, jz) if round(h_) == nivel]
            if g:
                print(f"  humano≈{nivel}: n={len(g)} · {juez} medio del juez = {st.mean(g):.2f}")
    print("\nInterpretación: kappa >= 0,6 se considera acuerdo sustancial. Un Spearman bajo entre humano y "
          "juez indica que la métrica automática no es confiable para este corpus.")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    m = sub.add_parser("muestra")
    m.add_argument("--results", required=True, type=Path)
    m.add_argument("--out", type=Path, default=Path("rag-eval/results/revision_humana.xlsx"))
    a = sub.add_parser("analizar")
    a.add_argument("--results", required=True, type=Path)
    a.add_argument("--revisor-a", required=True, type=Path)
    a.add_argument("--revisor-b", required=True, type=Path)
    args = ap.parse_args()
    (muestra if args.cmd == "muestra" else analizar)(args)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
