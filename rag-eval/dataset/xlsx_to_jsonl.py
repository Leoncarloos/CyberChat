"""Convierte la plantilla del conjunto de evaluación (XLSX) a JSONL y la valida.

Lee las hojas `Diccionario`, `Conjunto_80` y `Bateria`. No completa ni corrige nada:
solo informa lo que falta. Si hay errores no escribe los JSONL, salvo con
`--allow-incomplete`, que escribe únicamente las filas completas.

Uso:
    python rag-eval/dataset/xlsx_to_jsonl.py --xlsx ruta/plantilla.xlsx
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent

CONJUNTO_COLS = [
    "id", "tema_clave", "tema", "tipo", "organizacion", "historial_previo", "pregunta",
    "respuesta_referencia", "evidencia_esperada", "revisor_1", "revisor_2",
    "desacuerdo_resuelto", "notas",
]
BATERIA_COLS = [
    "id", "tipo", "tema_clave", "organizacion_consulta", "organizacion_contenido",
    "pregunta", "comportamiento_esperado", "revisor_1", "revisor_2", "notas",
]

POR_TEMA = {"documental": 8, "seguimiento": 2}
POR_TIPO_BATERIA = 10
ORG_ALIAS = re.compile(r"^ORG-[A-Z0-9]+$")
# Un turno previo por línea: «Usuario: ...» o «Asistente: ...».
TURNO = re.compile(r"^\s*(usuario|user|u|asistente|assistant|a)\s*:\s*(.*)$", re.IGNORECASE)
ROL = {"usuario": "user", "user": "user", "u": "user",
       "asistente": "assistant", "assistant": "assistant", "a": "assistant"}


def texto(valor) -> str:
    return "" if valor is None else str(valor).strip()


def leer_hoja(wb, nombre: str, columnas: list[str], errores: list[str]) -> list[dict]:
    if nombre not in wb.sheetnames:
        errores.append(f"Falta la hoja «{nombre}».")
        return []
    ws = wb[nombre]
    encabezado = [texto(c.value) for c in ws[1]]
    faltan = [c for c in columnas if c not in encabezado]
    if faltan:
        errores.append(f"{nombre}: faltan las columnas {', '.join(faltan)}.")
        return []
    idx = {c: encabezado.index(c) for c in columnas}
    filas = []
    for n, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not texto(fila[idx["id"]]):
            continue
        registro = {c: texto(fila[idx[c]]) for c in columnas}
        registro["_fila"] = n
        filas.append(registro)
    return filas


def leer_diccionario(wb, errores: list[str]) -> dict:
    if "Diccionario" not in wb.sheetnames:
        errores.append("Falta la hoja «Diccionario».")
        return {"temas": {}, "tipo_conjunto": set(), "tipo_bateria": set(), "revision": set()}
    ws = wb["Diccionario"]
    encabezado = [texto(c.value) for c in ws[1]]
    col = lambda nombre: encabezado.index(nombre)  # noqa: E731
    temas, tipo_conjunto, tipo_bateria, revision = {}, set(), set(), set()
    for fila in ws.iter_rows(min_row=2, values_only=True):
        clave = texto(fila[col("tema_clave")])
        if clave:
            temas[clave] = texto(fila[col("tema")])
        for destino, nombre in ((tipo_conjunto, "tipo_conjunto"), (tipo_bateria, "tipo_bateria"),
                                (revision, "revision")):
            valor = texto(fila[col(nombre)])
            if valor:
                destino.add(valor)
    return {"temas": temas, "tipo_conjunto": tipo_conjunto, "tipo_bateria": tipo_bateria,
            "revision": revision}


def historial_a_mensajes(crudo: str) -> tuple[list[dict], str]:
    """Convierte `historial_previo` en mensajes previos.

    Con prefijos «Usuario:» / «Asistente:» se respeta cada turno. Sin prefijos, todo el
    texto se toma como un único turno previo del usuario y se marca para revisión.
    """
    if not crudo:
        return [], "vacio"
    mensajes: list[dict] = []
    con_prefijo = False
    for linea in crudo.splitlines():
        if not linea.strip():
            continue
        m = TURNO.match(linea)
        if m:
            con_prefijo = True
            mensajes.append({"role": ROL[m.group(1).lower()], "content": m.group(2).strip()})
        elif mensajes:
            mensajes[-1]["content"] = f"{mensajes[-1]['content']}\n{linea.strip()}".strip()
        else:
            mensajes.append({"role": "user", "content": linea.strip()})
    return mensajes, ("turnos" if con_prefijo else "sin_prefijo")


def validar_conjunto(filas: list[dict], dic: dict, errores: list[str], avisos: list[str]):
    completas, por_tema = [], defaultdict(Counter)
    ids = Counter(f["id"] for f in filas)
    for dup in [i for i, n in ids.items() if n > 1]:
        errores.append(f"Conjunto_80: el id {dup} está repetido.")
    if len(filas) != 80:
        errores.append(f"Conjunto_80: se esperaban 80 consultas y hay {len(filas)}.")

    faltantes = defaultdict(list)
    for f in filas:
        ok = True
        if f["tema_clave"] not in dic["temas"]:
            errores.append(f"{f['id']}: tema_clave «{f['tema_clave']}» no está en el diccionario.")
            ok = False
        if f["tipo"] not in dic["tipo_conjunto"]:
            errores.append(f"{f['id']}: tipo «{f['tipo']}» no está en el diccionario.")
            ok = False
        else:
            por_tema[f["tema_clave"]][f["tipo"]] += 1
        for campo in ("pregunta", "respuesta_referencia", "organizacion"):
            if not f[campo]:
                faltantes[campo].append(f["id"])
                ok = False
        if f["organizacion"] and not ORG_ALIAS.match(f["organizacion"]):
            errores.append(f"{f['id']}: organizacion «{f['organizacion']}» no es un alias ORG-X.")
            ok = False
        if f["tipo"] == "seguimiento" and not f["historial_previo"]:
            faltantes["historial_previo (seguimiento)"].append(f["id"])
            ok = False
        if f["tipo"] == "documental" and f["historial_previo"]:
            avisos.append(f"{f['id']}: es «documental» pero tiene historial_previo; se ignora.")
        if not f["evidencia_esperada"]:
            faltantes["evidencia_esperada"].append(f["id"])
        for rev in ("revisor_1", "revisor_2"):
            if f[rev] and f[rev] not in dic["revision"]:
                errores.append(f"{f['id']}: {rev} «{f[rev]}» no es un valor válido.")
                ok = False
        if ok:
            completas.append(f)

    for campo, lista in faltantes.items():
        destino = avisos if campo == "evidencia_esperada" else errores
        destino.append(f"Conjunto_80: falta {campo} en {len(lista)} consultas ({resumir(lista)}).")
    for tema in dic["temas"]:
        for tipo, esperado in POR_TEMA.items():
            hay = por_tema[tema][tipo]
            if hay != esperado:
                errores.append(f"Conjunto_80: {tema} tiene {hay} «{tipo}» y deben ser {esperado}.")
    return completas


def validar_bateria(filas: list[dict], dic: dict, errores: list[str], avisos: list[str]):
    completas, por_tipo, faltantes = [], Counter(), defaultdict(list)
    ids = Counter(f["id"] for f in filas)
    for dup in [i for i, n in ids.items() if n > 1]:
        errores.append(f"Bateria: el id {dup} está repetido.")
    for f in filas:
        ok = True
        if f["tipo"] not in dic["tipo_bateria"]:
            errores.append(f"{f['id']}: tipo «{f['tipo']}» no está en el diccionario.")
            ok = False
        else:
            por_tipo[f["tipo"]] += 1
        if f["tema_clave"] and f["tema_clave"] not in dic["temas"]:
            errores.append(f"{f['id']}: tema_clave «{f['tema_clave']}» no está en el diccionario.")
            ok = False
        for campo in ("pregunta", "organizacion_consulta", "comportamiento_esperado"):
            if not f[campo]:
                faltantes[campo].append(f["id"])
                ok = False
        for campo in ("organizacion_consulta", "organizacion_contenido"):
            if f[campo] and not ORG_ALIAS.match(f[campo]):
                errores.append(f"{f['id']}: {campo} «{f[campo]}» no es un alias ORG-X.")
                ok = False
        if f["tipo"] == "aislamiento":
            if not f["organizacion_contenido"]:
                faltantes["organizacion_contenido (aislamiento)"].append(f["id"])
                ok = False
            elif f["organizacion_contenido"] == f["organizacion_consulta"]:
                errores.append(f"{f['id']}: en «aislamiento» la organización que consulta y la "
                               "dueña del contenido deben ser distintas.")
                ok = False
        for rev in ("revisor_1", "revisor_2"):
            if f[rev] and f[rev] not in dic["revision"]:
                errores.append(f"{f['id']}: {rev} «{f[rev]}» no es un valor válido.")
                ok = False
        if ok:
            completas.append(f)
    for campo, lista in faltantes.items():
        errores.append(f"Bateria: falta {campo} en {len(lista)} casos ({resumir(lista)}).")
    for tipo in sorted(dic["tipo_bateria"]):
        if por_tipo[tipo] != POR_TIPO_BATERIA:
            errores.append(f"Bateria: hay {por_tipo[tipo]} casos «{tipo}» y deben ser "
                           f"{POR_TIPO_BATERIA}.")
    return completas


def resumir(ids: list[str], maximo: int = 8) -> str:
    if len(ids) <= maximo:
        return ", ".join(ids)
    return f"{', '.join(ids[:maximo])} y {len(ids) - maximo} más"


def registro_conjunto(f: dict, dic: dict, avisos: list[str]) -> dict:
    historial, modo = ([], "vacio")
    if f["tipo"] == "seguimiento":
        historial, modo = historial_a_mensajes(f["historial_previo"])
        if modo == "sin_prefijo":
            avisos.append(f"{f['id']}: historial_previo sin prefijos «Usuario:»/«Asistente:»; "
                          "se toma como un solo turno del usuario.")
    return {
        "id": f["id"],
        "tema_clave": f["tema_clave"],
        "tema": dic["temas"][f["tema_clave"]],
        "tipo": f["tipo"],
        "org_alias": f["organizacion"],
        "history": historial,
        "history_parse": modo,
        "user_input": f["pregunta"],
        "reference": f["respuesta_referencia"],
        "expected_evidence": f["evidencia_esperada"],
        "review": {"revisor_1": f["revisor_1"], "revisor_2": f["revisor_2"],
                   "desacuerdo_resuelto": f["desacuerdo_resuelto"]},
        "notas": f["notas"],
    }


def registro_bateria(f: dict, dic: dict) -> dict:
    return {
        "id": f["id"],
        "tipo": f["tipo"],
        "tema_clave": f["tema_clave"],
        "tema": dic["temas"].get(f["tema_clave"], ""),
        "org_alias": f["organizacion_consulta"],
        "org_contenido": f["organizacion_contenido"],
        "user_input": f["pregunta"],
        "expected_behavior": f["comportamiento_esperado"],
        "review": {"revisor_1": f["revisor_1"], "revisor_2": f["revisor_2"]},
        "notas": f["notas"],
    }


def escribir_jsonl(ruta: Path, registros: list[dict]) -> str:
    contenido = "".join(json.dumps(r, ensure_ascii=False, sort_keys=True) + "\n" for r in registros)
    ruta.write_text(contenido, encoding="utf-8", newline="\n")
    return hashlib.sha256(contenido.encode("utf-8")).hexdigest()


def revision(filas: list[dict]) -> dict:
    ambos = sum(1 for f in filas if f["revisor_1"] == "Sí" and f["revisor_2"] == "Sí")
    alguno = sum(1 for f in filas if f["revisor_1"] or f["revisor_2"])
    return {"validadas_por_ambos": ambos, "con_alguna_revision": alguno, "total": len(filas)}


def main() -> int:
    for flujo in (sys.stdout, sys.stderr):
        if hasattr(flujo, "reconfigure"):
            flujo.reconfigure(encoding="utf-8")

    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--xlsx", required=True, type=Path, help="Plantilla completada.")
    ap.add_argument("--out", type=Path, default=HERE, help="Carpeta de salida (por defecto, esta).")
    ap.add_argument("--allow-incomplete", action="store_true",
                    help="Escribe solo las filas completas aunque haya errores.")
    args = ap.parse_args()

    if not args.xlsx.exists():
        print(f"No existe el archivo {args.xlsx}", file=sys.stderr)
        return 2

    errores: list[str] = []
    avisos: list[str] = []
    wb = load_workbook(args.xlsx, data_only=False)
    dic = leer_diccionario(wb, errores)
    conjunto = leer_hoja(wb, "Conjunto_80", CONJUNTO_COLS, errores)
    bateria = leer_hoja(wb, "Bateria", BATERIA_COLS, errores)

    conjunto_ok = validar_conjunto(conjunto, dic, errores, avisos)
    bateria_ok = validar_bateria(bateria, dic, errores, avisos)

    rev_conjunto = revision(conjunto)
    if conjunto and rev_conjunto["validadas_por_ambos"] < len(conjunto):
        avisos.append(f"Conjunto_80: solo {rev_conjunto['validadas_por_ambos']} de {len(conjunto)} "
                      "referencias están validadas por ambos revisores.")

    escribe = not errores or args.allow_incomplete
    salida = {}
    if escribe:
        args.out.mkdir(parents=True, exist_ok=True)
        reg_conjunto = [registro_conjunto(f, dic, avisos) for f in conjunto_ok]
        reg_bateria = [registro_bateria(f, dic) for f in bateria_ok]
        salida = {
            "conjunto.jsonl": {"registros": len(reg_conjunto),
                               "sha256": escribir_jsonl(args.out / "conjunto.jsonl", reg_conjunto)},
            "bateria.jsonl": {"registros": len(reg_bateria),
                              "sha256": escribir_jsonl(args.out / "bateria.jsonl", reg_bateria)},
        }

    informe = {
        "generado": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "xlsx": args.xlsx.name,
        "xlsx_sha256": hashlib.sha256(args.xlsx.read_bytes()).hexdigest(),
        "valido": not errores,
        "conjunto": {"filas": len(conjunto), "completas": len(conjunto_ok), "revision": rev_conjunto,
                     "organizaciones": sorted({f["organizacion"] for f in conjunto if f["organizacion"]})},
        "bateria": {"filas": len(bateria), "completas": len(bateria_ok),
                    "revision": revision(bateria)},
        "salida": salida,
        "errores": errores,
        "avisos": avisos,
    }
    args.out.mkdir(parents=True, exist_ok=True)
    (args.out / "validacion.json").write_text(
        json.dumps(informe, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")

    print(f"Conjunto_80: {len(conjunto_ok)} de {len(conjunto)} consultas completas")
    print(f"Bateria:     {len(bateria_ok)} de {len(bateria)} casos completos")
    for titulo, lista in (("ERRORES", errores), ("AVISOS", avisos)):
        if lista:
            print(f"\n{titulo} ({len(lista)})")
            for linea in lista:
                print(f"  - {linea}")
    if salida:
        print("\nEscrito:")
        for nombre, datos in salida.items():
            print(f"  {args.out / nombre}  ({datos['registros']} registros, sha256 {datos['sha256'][:12]})")
    else:
        print("\nNo se escribieron los JSONL: corrige los errores o usa --allow-incomplete.")
    print(f"Informe: {args.out / 'validacion.json'}")
    return 0 if not errores else 1


if __name__ == "__main__":
    sys.exit(main())
